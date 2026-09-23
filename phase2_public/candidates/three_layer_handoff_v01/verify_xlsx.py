"""Independent read-only XLSX verification. Does not save/edit the workbook."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook


def verify(output_dir):
    path = output_dir / "三层协议_离线复核.xlsx"
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    packet = json.loads((output_dir / "workbook_packet.json").read_text(encoding="utf-8"))
    wb = load_workbook(path, data_only=False)
    if wb.sheetnames != [s["name"] for s in packet["sheets"]]:
        raise AssertionError("sheet_set_or_order")
    summary = []
    literal_count = 0
    for s in packet["sheets"]:
        ws = wb[s["name"]]
        expected = [s["columns"], *s["rows"]]
        for i, row in enumerate(expected, 6):
            for j, value in enumerate(row, 1):
                cell = ws.cell(i, j)
                if cell.data_type in ("f", "e"):
                    raise AssertionError(f"unexpected_formula_or_error:{ws.title}:{cell.coordinate}")
                if cell.value != value or (isinstance(value, bool) and cell.data_type != "b"):
                    raise AssertionError(f"value_type_mismatch:{ws.title}:{cell.coordinate}:{cell.value!r}!={value!r}")
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
                    literal_count += 1
                    if cell.data_type != "s": raise AssertionError("unsafe_literal_text")
        if ws.freeze_panes != "C7" or ws.sheet_view.showGridLines is not False:
            raise AssertionError("pane_or_grid_setting_lost")
        if len(ws.tables) != (1 if s["rows"] else 0): raise AssertionError("table_filter_missing")
        summary.append({"sheet":s["name"], "data_rows":len(s["rows"]), "cells_checked":len(expected)*len(s["columns"]),
                        "freeze_panes":ws.freeze_panes, "tables":len(ws.tables)})
    main = wb[packet["sheets"][0]["name"]]
    if len(main.conditional_formatting) != 2: raise AssertionError("conditional_formats_missing")
    if hashlib.sha256(path.read_bytes()).hexdigest() != before: raise AssertionError("workbook_modified")
    return {"status":"passed", "reader":"openpyxl_read_only_no_save", "sheets":summary,
            "formula_like_literal_cells_verified":literal_count, "workbook_hash_unchanged":True,
            "workbook_sha256":before, "source_values_and_nulls_preserved":True,
            "desktop_excel_application_tested":False}


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--output", type=Path, required=True)
    args = p.parse_args()
    result = verify(args.output)
    with (args.output / "independent_readback.json").open("x", encoding="utf-8") as out:
        json.dump(result, out, ensure_ascii=False, indent=2)
        out.write("\n")
    print(json.dumps(result, ensure_ascii=False))
