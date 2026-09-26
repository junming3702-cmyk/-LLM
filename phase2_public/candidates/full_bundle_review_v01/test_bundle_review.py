"""Synthetic development checks only; no real project, law or API traffic."""

from __future__ import annotations

import json
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfWriter

from bundle_review import find_pair, intake_verified_pagewise_cache, pair_score, require_manifest, run
from assemble import assemble, excel_records
from export_review_excel import build_workbook, flatten_record
from pair_reasoner import apply_pair_gate, normalize_pair_shape, replay_stored_result


TENDER_QUALIFICATION = "3.5.5 投标人须提供有效的建筑工程施工总承包资质证书。"
TENDER_GUARANTEE = "投标保证金金额不得超过项目估算价的2%。"
BID_QUALIFICATION = "3.5.5 我方提供有效的建筑工程施工总承包资质证书。"
BID_GUARANTEE = "投标保证金金额按项目估算价的2%提供。"


def make_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    document.save(path)


class BundleDevelopmentTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        make_docx(self.root / "tender.docx", [TENDER_QUALIFICATION, TENDER_GUARANTEE, "项目概况。"])
        make_docx(self.root / "bid.docx", [BID_QUALIFICATION, BID_GUARANTEE, "公司简介。"])
        self.manifest = {
            "project_id": "SYN-BUNDLE-001", "project_type": "construction_project",
            "project_location": {"province": "四川省", "human_confirmation": "unconfirmed"},
            "documents": [
                {"document_key": "T1", "role": "tender", "path": "tender.docx", "declared_issued": True, "source_status": "synthetic"},
                {"document_key": "B1", "role": "final_bid", "path": "bid.docx", "declared_final_submitted": True, "source_status": "synthetic"},
            ],
        }
        self.manifest_path = self.root / "bundle_manifest.json"
        self.manifest_path.write_text(json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8")

    def test_manifest_rejects_unfinal_bid_and_award_material(self) -> None:
        self.manifest["documents"][1]["declared_final_submitted"] = False
        self.manifest_path.write_text(json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "final_bid_submission_declaration_required"):
            require_manifest(self.manifest_path)
        self.manifest["documents"][1]["role"] = "award_record"
        self.manifest_path.write_text(json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "only_explicit_tender_and_final_bid_roles_allowed"):
            require_manifest(self.manifest_path)

    def test_candidate_discovery_pairing_and_core_labels(self) -> None:
        out = self.root / "run"
        result = run(self.manifest_path, out)
        self.assertEqual(result["documents"], 2)
        issues = json.loads((out / "candidates.json").read_text(encoding="utf-8"))
        coverage = json.loads((out / "coverage.json").read_text(encoding="utf-8"))
        labels = [json.loads(line) for line in (out / "candidate_labels.jsonl").read_text(encoding="utf-8").splitlines()]
        tasks = {row["task"] for row in issues}
        self.assertEqual(tasks, {"tender_clause_legality", "bid_responsiveness", "bid_standalone_legality"})
        qualification = next(row for row in issues if row["task"] == "bid_responsiveness" and TENDER_QUALIFICATION in row["primary_source"]["quote"])
        self.assertEqual(qualification["pairing"]["status"], "provisional_match_needs_verification")
        self.assertIn(BID_QUALIFICATION, qualification["pairing"]["candidate_matches"][0]["source"]["quote"])
        # Fixed synthetic annotation: two tender requirements and their paired
        # final-bid paragraphs. This is development-fixture coverage, not an
        # independent estimate of document-wide recall or legal accuracy.
        expected_pairs = {TENDER_QUALIFICATION: BID_QUALIFICATION, TENDER_GUARANTEE: BID_GUARANTEE}
        paired_rows = [row for row in issues if row["task"] == "bid_responsiveness"]
        discovered = sum(any(t in row["primary_source"]["quote"] for row in paired_rows) for t in expected_pairs)
        paired = sum(any(t in row["primary_source"]["quote"] and row["pairing"]["status"] == "provisional_match_needs_verification"
                         and b in row["pairing"]["candidate_matches"][0]["source"]["quote"] for row in paired_rows)
                     for t, b in expected_pairs.items())
        self.assertEqual((discovered, len(expected_pairs)), (2, 2))
        self.assertEqual((paired, len(expected_pairs)), (2, 2))
        self.assertEqual(coverage["independent_item_recall"], None)
        self.assertEqual(coverage["documents"][0]["page_locator_limitation"], "DOCX has structural paragraph/table locators only")
        self.assertTrue(all(label["bundle_evidence"]["documents_received"] for label in labels))
        assembled = assemble(out)
        self.assertEqual(assembled["core_result_count"], 0)
        self.assertEqual(assembled["three_layer_valid_count"], 0)
        self.assertTrue(all(row["core_status"] in {"not_run", "not_eligible_unpaired"} for row in assembled["records"]))
        workbook_rows = [row for rec in excel_records(assembled) for row in flatten_record(rec)]
        self.assertGreaterEqual(len(workbook_rows), len(issues))
        self.assertTrue(any("招标依据" in row[2] and "投标原文" in row[2] for row in workbook_rows))
        self.assertTrue(all(row[10] == "requires_human_second_review" for row in workbook_rows))
        workbook_path = self.root / "synthetic_review.xlsx"
        build_workbook(excel_records(assembled), workbook_path)
        workbook = load_workbook(workbook_path)
        self.assertIn("Human review", workbook.sheetnames)
        self.assertTrue(any("物理页码未建立" in str(row[2].value or "")
                            for row in workbook["Human review"].iter_rows(min_row=2)))

    def test_pairing_does_not_equate_no_match_with_missing_bid(self) -> None:
        self.assertLess(pair_score("防火等级须达到甲级。", "公司地址在甲城市。"), 0.18)
        self.assertEqual(pair_score("", "响应"), 0.0)
        requirement = {"quote": "应提供资格证书。"}
        candidates = [{"quote": "提供资格证书。", "document_key": "B1", "block_id": "a"},
                      {"quote": "提供资格证书。", "document_key": "B1", "block_id": "b"}]
        self.assertEqual(find_pair(requirement, candidates)["status"], "ambiguous_needs_human_matching")

    def test_same_section_number_without_textual_alignment_cannot_pair(self) -> None:
        requirement = {"quote": "3.2.1 投标人应按第五章工程量清单的要求填写相应表格。"}
        unrelated = {"quote": "3.2.1 火灾报警自检功能。", "document_key": "B1", "block_id": "fire"}
        self.assertLess(pair_score(requirement["quote"], unrelated["quote"]), 0.18)
        self.assertEqual(find_pair(requirement, [unrelated])["status"], "not_found_in_reviewed_text")

    def test_pagewise_cache_requires_exact_source_and_artifact_hashes(self) -> None:
        source = self.root / "source.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(source)
        file_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        blocks = self.root / "blocks.json"
        blocks.write_text(json.dumps([{"document_id": "X", "source_sha256": file_hash,
                                       "block_id": "X-p0001-b001", "page": 1, "type": "table",
                                       "text": "我方承诺满足招标条件。", "native_locator_valid": True}], ensure_ascii=False), encoding="utf-8")
        preflight = self.root / "preflight.json"
        preflight.write_text(json.dumps({"sha256": file_hash, "enumerated_page_count": 1,
                                         "verdict": "UNAVAILABLE", "warnings": ["xref advisory"]}), encoding="utf-8")
        cache = {"document_alias": "X", "source_sha256": file_hash,
                 "pagewise_blocks_path": str(blocks), "preflight_path": str(preflight),
                 "blocks_sha256": hashlib.sha256(blocks.read_bytes()).hexdigest(),
                 "preflight_sha256": hashlib.sha256(preflight.read_bytes()).hexdigest()}
        declared = {"resolved_path": str(source), "document_key": "X", "role": "final_bid",
                    "verified_pagewise_cache": cache}
        result = intake_verified_pagewise_cache(declared, "SYN", self.root / "cached")
        quality = json.loads(Path(result["quality_report_path"]).read_text(encoding="utf-8"))
        self.assertEqual(quality["quality_status"], "needs_human_review")
        self.assertEqual(quality["empty_pages"], [])
        self.assertFalse(quality["ocr_semantic_accuracy_verified"])
        self.assertEqual(len(Path(quality["page_ledger_path"]).read_text(encoding="utf-8").splitlines()), 1)
        manifest = {"project_id": "SYN-CACHE", "project_type": "construction_project", "documents": [
            {"document_key": "T1", "role": "tender", "path": "tender.docx", "declared_issued": True},
            {"document_key": "X", "role": "final_bid", "path": "source.pdf", "declared_final_submitted": True,
             "verified_pagewise_cache": cache}]}
        manifest_path = self.root / "cached_manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        run(manifest_path, self.root / "cached_run")
        candidates = json.loads((self.root / "cached_run" / "candidates.json").read_text(encoding="utf-8"))
        self.assertTrue(any(c["task"] == "bid_standalone_legality" and c["primary_source"]["block_type"] == "table"
                            for c in candidates))
        cache["source_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "cached_source_hash_mismatch"):
            intake_verified_pagewise_cache(declared, "SYN", self.root / "tampered")

    def test_paired_text_gate_does_not_invent_legal_basis(self) -> None:
        out = self.root / "run"
        run(self.manifest_path, out)
        label = json.loads((out / "pair_labels.jsonl").read_text(encoding="utf-8").splitlines()[0])
        primary = label["bundle_evidence"]["primary"]
        paired = label["bundle_evidence"]["paired_bid_source"]
        raw = {"findings": [{"issue_id": label["issue_id"], "tender_quote": primary["quote"],
                             "bid_quote": paired["quote"], "difference": "证书名称需人工核对。",
                             "status": "potential_nonresponse", "recommended_human_action": "核对资质证书原件。"}]}
        result = apply_pair_gate(raw, label)
        self.assertFalse(result["blocked"])
        self.assertEqual(result["response"]["findings"][0]["legal_evidence"], [])
        self.assertEqual(result["response"]["findings"][0]["conclusion_type"], "potential_risk")
        normalized, action = normalize_pair_shape(raw["findings"][0])
        self.assertEqual(action, "wrapped_exact_single_finding_root")
        self.assertEqual(normalized, raw)
        self.assertEqual(normalize_pair_shape({"issue_id": label["issue_id"]})[1], None)
        altered = json.loads(json.dumps(raw, ensure_ascii=False))
        altered["findings"][0]["bid_quote"] = "我方具备其他证书。"
        self.assertEqual(apply_pair_gate(altered, label)["status"], "blocked")
        self.assertEqual(apply_pair_gate(raw, label, finish_reason="length")["status"], "blocked")
        pair_dir = self.root / "pair"
        pair_dir.mkdir()
        pair_result = {"issue_id": label["issue_id"], "runtime_input": {
            "review_task_kind": "bid_responsiveness", "bundle_evidence": label["bundle_evidence"],
            "contract_evidence": {"document_excerpt": label["document_excerpt"]},
            "review_scope": {"documents_received": label["bundle_evidence"]["documents_received"]}},
            "final_llm_response": raw, "post_llm_gate": result}
        (pair_dir / f"{label['issue_id']}.json").write_text(json.dumps(pair_result, ensure_ascii=False), encoding="utf-8")
        flat_prior = {**pair_result, "final_llm_response": raw["findings"][0],
                      "provider_diagnostics": {"finish_reason": "stop", "ok": True},
                      "post_llm_gate": apply_pair_gate(raw["findings"][0], label)}
        prior_path = self.root / "flat_prior.json"
        prior_path.write_text(json.dumps(flat_prior, ensure_ascii=False), encoding="utf-8")
        replayed = replay_stored_result(label, prior_path)
        self.assertFalse(replayed["post_llm_gate"]["blocked"])
        self.assertTrue(replayed["offline_replay_no_api_call"])
        self.assertEqual(replayed["final_llm_response"], raw["findings"][0])
        assembled = assemble(out, pair_dir=pair_dir)
        row = next(x for x in assembled["records"] if x["issue_id"] == label["issue_id"])
        self.assertEqual(row["presentation_response"]["findings"][0]["conclusion_type"], "potential_risk")

    def test_empty_pdf_pages_remain_visible_and_not_clean(self) -> None:
        pdf = self.root / "blank_bid.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(pdf)
        self.manifest["documents"][1]["path"] = pdf.name
        self.manifest_path.write_text(json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8")
        out = self.root / "blank-run"
        run(self.manifest_path, out)
        assembled = assemble(out)
        bid_coverage = next(row for row in assembled["coverage"]["documents"] if row["document_key"] == "B1")
        self.assertEqual(bid_coverage["unreadable_or_empty_physical_pages"], [1])
        self.assertEqual(bid_coverage["conditional_enhancement_recommended"], "MinerU_then_local_coordinate_OCR_if_unresolved")
        rows = [row for rec in excel_records(assembled) for row in flatten_record(rec)]
        self.assertTrue(any("未成功识别页" in row[6] for row in rows))
        candidates = json.loads((out / "candidates.json").read_text(encoding="utf-8"))
        self.assertTrue(all(row["pairing"]["status"] == "not_found_in_reviewed_text"
                            for row in candidates if row["task"] == "bid_responsiveness"))
        labels = [json.loads(line) for line in (out / "candidate_labels.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertFalse(any(row["review_task_kind"] == "bid_responsiveness" for row in labels))

    def test_core_binding_and_three_layer_handoff_are_supplementary(self) -> None:
        out = self.root / "run"
        run(self.manifest_path, out)
        issues = json.loads((out / "candidates.json").read_text(encoding="utf-8"))
        issue = next(x for x in issues if x["task"] == "tender_clause_legality")
        labels = {r["issue_id"]: r for r in (
            json.loads(line) for line in (out / "candidate_labels.jsonl").read_text(encoding="utf-8").splitlines()
        )}
        label = labels[issue["issue_id"]]
        core_dir = self.root / "core"
        core_dir.mkdir()
        gated = {"status": "review_required", "blocked": False,
                 "response": {"findings": [{"issue_id": issue["issue_id"],
                               "conclusion_type": "requires_human_legal_review",
                               "legal_evidence": [], "document_excerpt": label["document_excerpt"],
                               "assistant_recommendation": "测试风险候选，人工复核。"}]}}
        core = {"issue_id": issue["issue_id"], "runtime_input": {
            "review_task_kind": issue["task"], "bundle_evidence": label["bundle_evidence"],
            "contract_evidence": {"document_excerpt": label["document_excerpt"]},
            "review_scope": {"documents_received": label["bundle_evidence"]["documents_received"]}},
            "post_llm_gate": gated, "final_llm_response": {"synthetic_test_only": True}}
        (core_dir / f"{issue['issue_id']}.json").write_text(json.dumps(core, ensure_ascii=False), encoding="utf-8")

        protocol_dir = Path(__file__).resolve().parents[1] / "three_layer_handoff_v01"
        sys.path.insert(0, str(protocol_dir))
        import fixtures  # noqa: PLC0415
        from gate import seal_context  # noqa: PLC0415
        ctx = fixtures.context()
        src = issue["primary_source"]
        ctx["sources"] = [{"source_id": "S1", "document_sha256": src["document_sha256"],
                           "locator": src["source_locator"], "text": src["quote"]}]
        ctx["evidence_registry"] = [{"evidence_id": "E1", "source_id": "S1",
                                      "document_sha256": src["document_sha256"],
                                      "locator": src["source_locator"], "quote": src["quote"]}]
        ctx = seal_context(ctx)
        response = fixtures.response(ctx)
        response["claim_assessments"] = [fixtures.claim("C1", checks=["CK1"])]
        response["completed_checks"] = [fixtures.check(refs=["E1"])]
        handoff_path = self.root / "handoff.json"
        handoff_path.write_text(json.dumps([{"issue_id": issue["issue_id"], "context": ctx,
            "raw_response": response, "execution": {"finish_reason": "stop", "transport_ok": True}}], ensure_ascii=False), encoding="utf-8")
        assembled = assemble(out, core_dir, handoff_path)
        self.assertEqual(assembled["core_result_count"], 1)
        self.assertEqual(assembled["three_layer_valid_count"], 1)
        row = next(x for x in assembled["records"] if x["issue_id"] == issue["issue_id"])
        self.assertEqual(row["presentation_response"]["findings"][0]["conclusion_type"], "requires_human_legal_review")
        self.assertEqual(row["risk_response_gated_unmodified"], gated)
        self.assertEqual(row["three_layer_handoff"]["handoff_status"], "ready_for_complete_review")
        self.assertTrue(row["three_layer_handoff"]["presentation_only"])
        embedded = dict(core)
        embedded["three_layer_protocol"] = json.loads(handoff_path.read_text(encoding="utf-8"))[0]
        (core_dir / f"{issue['issue_id']}.json").write_text(json.dumps(embedded, ensure_ascii=False), encoding="utf-8")
        embedded_result = assemble(out, core_dir)
        embedded_row = next(x for x in embedded_result["records"] if x["issue_id"] == issue["issue_id"])
        self.assertEqual(embedded_row["three_layer_handoff"]["processing_status"], "valid")
        self.assertEqual(embedded_row["presentation_response"]["findings"][0]["conclusion_type"], "requires_human_legal_review")
        with self.assertRaisesRegex(ValueError, "duplicate_three_layer_source_for_issue"):
            assemble(out, core_dir, handoff_path)
        tampered = dict(core)
        tampered["runtime_input"] = {**core["runtime_input"], "review_task_kind": "bid_responsiveness"}
        (core_dir / f"{issue['issue_id']}.json").write_text(json.dumps(tampered, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "core_result_bundle_binding_mismatch"):
            assemble(out, core_dir, handoff_path)


if __name__ == "__main__":
    unittest.main()
